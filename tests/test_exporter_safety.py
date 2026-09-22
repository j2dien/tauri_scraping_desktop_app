import sys
import tempfile
import types
import unittest
from pathlib import Path


# openpyxl adalah dependency runtime aplikasi, tetapi venv test minimal tidak
# selalu memasangnya. Gunakan implementasi asli bila tersedia dan stub hanya
# sebagai fallback untuk test helper murni.
try:
    import openpyxl
    from openpyxl import load_workbook

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    openpyxl = types.ModuleType("openpyxl")
    openpyxl.__path__ = []
    openpyxl.Workbook = type("Workbook", (), {})

    styles = types.ModuleType("openpyxl.styles")
    for name in ("Font", "PatternFill", "Alignment", "Border", "Side"):
        setattr(styles, name, type(name, (), {"__init__": lambda self, *args, **kwargs: None}))

    utils = types.ModuleType("openpyxl.utils")
    utils.get_column_letter = lambda value: str(value)

    sys.modules["openpyxl"] = openpyxl
    sys.modules["openpyxl.styles"] = styles
    sys.modules["openpyxl.utils"] = utils


from core.exporter import (
    INSTAGRAM_DETAIL_HEADERS,
    _optional_count_for_excel,
    clean_cell_value,
    export_to_excel,
)


class ExcelCellSafetyTests(unittest.TestCase):
    def test_instagram_detail_headers_match_requested_schema(self):
        self.assertEqual(
            INSTAGRAM_DETAIL_HEADERS,
            [
                "Username",
                "Teks Komentar",
                "Sudah Like Post?",
                "Like Komentar",
                "Tanggal Komentar",
                "Post URL",
                "Like Postingan",
                "Tanggal Post",
                "Caption Post",
            ],
        )

    def test_formula_prefixes_are_exported_as_plain_text(self):
        for value in (
            "=HYPERLINK(\"https://example.invalid\")",
            "+1+1",
            "-2+3",
            "@SUM(A1:A2)",
            "   =1+1",
        ):
            with self.subTest(value=value):
                self.assertTrue(clean_cell_value(value).startswith("'"))

    def test_normal_text_is_unchanged_and_control_characters_are_removed(self):
        self.assertEqual(clean_cell_value("komentar biasa"), "komentar biasa")
        self.assertEqual(clean_cell_value("abc\x00def"), "abcdef")

    def test_optional_count_helper_also_sanitizes_untrusted_strings(self):
        self.assertEqual(_optional_count_for_excel(None), "Tidak tersedia")
        self.assertEqual(_optional_count_for_excel("=1+1"), "'=1+1")

    @unittest.skipUnless(OPENPYXL_AVAILABLE, "openpyxl runtime tidak tersedia")
    def test_instagram_workbook_uses_all_comments_and_exact_nine_columns(self):
        full_caption = "Caption lengkap " + ("panjang " * 30)
        all_comments = [
            {
                "commenter_username": "all_comments_user",
                "comment_text": "komentar",
                "has_liked_post": "Belum dapat diverifikasi",
                "comment_likes": 0,
                "comment_date": "2026-09-20 10:00:00",
                "post_url": "https://www.instagram.com/p/ABC/",
                "post_likes": 12,
                "post_date": "2026-09-20 09:00:00",
                "post_caption": full_caption,
            }
        ]
        detail_comments = [{"commenter_username": "top_n_only"}]

        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / "instagram.xlsx"
            exported = export_to_excel(
                top_commenters=[],
                detail_comments=detail_comments,
                all_comments=all_comments,
                scraped_posts=[],
                summary_stats={},
                target_username="target",
                start_date="20-09-2026",
                end_date="20-09-2026",
                platform="Instagram",
                filename=str(target),
            )
            workbook = load_workbook(exported, read_only=True)
            sheet = workbook["Detail Komentar"]

            self.assertEqual(
                [cell.value for cell in sheet[1]],
                INSTAGRAM_DETAIL_HEADERS,
            )
            self.assertEqual(sheet.max_column, 9)
            self.assertEqual(sheet.cell(row=2, column=1).value, "all_comments_user")
            self.assertEqual(sheet.cell(row=2, column=9).value, full_caption)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
